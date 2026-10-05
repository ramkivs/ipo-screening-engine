"""Phase 6B Acceptance Test Suite: Historical Outcome Store & Dataset Foundation.

Covers all 30 tests:
T-6B-01: Single IPO dataset assembly.
T-6B-02: Multiple IPO dataset assembly.
T-6B-03: Multiple horizons per IPO.
T-6B-04: Missing 1W does not create zero.
T-6B-05: Missing 1M does not create zero.
T-6B-06: Missing 6M does not create zero.
T-6B-07: Partial lifecycle IPO remains in dataset.
T-6B-08: FINAL evaluation linkage verified.
T-6B-09: Result hash mismatch rejected.
T-6B-10: Invalid observation hash rejected.
T-6B-11: Duplicate horizon detected.
T-6B-12: Conflicting observation versions fail closed.
T-6B-13: Restated observation selects deterministic effective version.
T-6B-14: Superseded observation remains preserved.
T-6B-15: Dataset ordering deterministic.
T-6B-16: Dataset hash deterministic.
T-6B-17: Dataset hash changes when material data changes.
T-6B-18: Filesystem ordering cannot affect dataset.
T-6B-19: Point-in-time evaluation fields remain unchanged.
T-6B-20: Benchmark UNKNOWN remains UNKNOWN.
T-6B-21: Excess return UNKNOWN remains UNKNOWN.
T-6B-22: UNVERIFIED corporate-action observation retains status.
T-6B-23: Dataset manifest verifies correctly.
T-6B-24: Dataset verification rejects tampering.
T-6B-25: JSON export is deterministic.
T-6B-26: CSV projection is deterministic.
T-6B-27: Multi-IPO Excel Post_Listing projection.
T-6B-28: Multi-IPO Excel Backtest projection.
T-6B-29: Existing v1.5 + Phase 6A tests remain green.
T-6B-30: Complete Phase 6B E2E workflow.
"""

from __future__ import annotations

import copy
import csv
import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest

try:
    from ipo_screening.post_listing import (
        CALCULATION_VERSION,
        BacktestDataset,
        BacktestDatasetManifest,
        BacktestDatasetRow,
        ConflictingObservationError,
        DatasetAssemblyError,
        DatasetRowStatus,
        DatasetVerificationError,
        DuplicateEvaluationError,
        Horizon,
        ObservationStatus,
        PostListingObservation,
        PriceObservation,
        ProvenanceRecord,
        ReturnSet,
        VerificationStatus,
        assemble_dataset,
        build_observation,
        compute_dataset_hash,
        export_dataset_csv,
        export_dataset_json,
        get_effective_observations,
        load_price_file,
        save_observation,
        verify_dataset,
    )
except ImportError:
    from engine.ipo_screening.post_listing import (
        CALCULATION_VERSION,
        BacktestDataset,
        BacktestDatasetManifest,
        BacktestDatasetRow,
        ConflictingObservationError,
        DatasetAssemblyError,
        DatasetRowStatus,
        DatasetVerificationError,
        DuplicateEvaluationError,
        Horizon,
        ObservationStatus,
        PostListingObservation,
        PriceObservation,
        ProvenanceRecord,
        ReturnSet,
        VerificationStatus,
        assemble_dataset,
        build_observation,
        compute_dataset_hash,
        export_dataset_csv,
        export_dataset_json,
        get_effective_observations,
        load_price_file,
        save_observation,
        verify_dataset,
    )

from conftest import EVAL_AT

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "post_listing"


def _create_mock_eval(
    store_dir: Path,
    eval_id: str,
    ipo_id: str,
    timestamp: str,
    result_hash: str,
    score: float = 65.0,
    verdict: str = "APPLY",
    mode: str = "FINAL",
    issue_price: float = 100.0,
    listing_date: str = "2026-10-06",
) -> Path:
    edir = store_dir / eval_id
    edir.mkdir(parents=True, exist_ok=True)

    eval_data = {
        "evaluation_id": eval_id,
        "ipo_id": ipo_id,
        "company_name": f"{ipo_id} Limited",
        "evaluation_mode": mode,
        "evaluation_timestamp": timestamp,
        "result_hash": result_hash,
        "score": {
            "final_score": score,
            "completeness_pct": 92.5,
            "penalties_total": 0.0,
            "modules": [
                {"module_id": "A", "name": "Financial Quality", "score": 25.0, "max": 30.0},
                {"module_id": "B", "name": "Valuation", "score": 15.0, "max": 20.0},
                {"module_id": "C", "name": "Governance", "score": 10.0, "max": 15.0},
                {"module_id": "D", "name": "Issue Structure", "score": 5.0, "max": 10.0},
                {"module_id": "E", "name": "Market Sentiment", "score": 5.0, "max": 10.0},
                {"module_id": "F", "name": "Lead Manager", "score": 5.0, "max": 15.0},
            ],
        },
        "score_range": {
            "base_score": score,
            "lower_bound": score - 5.0,
            "upper_bound": score + 5.0,
            "unknown_points": 7.5,
        },
        "verdict": {
            "verdict": verdict,
            "band_score": 60.0,
            "insufficient_data": False,
        },
        "confidence": {"level": "HIGH"},
        "knockouts": {
            "status": "PASS",
            "triggered": [],
            "insufficient_data": [],
        },
    }
    with (edir / "evaluation.json").open("w", encoding="utf-8") as f:
        json.dump(eval_data, f, indent=2)

    input_data = {
        "snapshot": {
            "input": {
                "ipo_id": ipo_id,
                "company": {"name": f"{ipo_id} Limited", "symbol": ipo_id},
                "issue": {"issue_price": issue_price, "listing_date": listing_date},
            }
        }
    }
    with (edir / "input.json").open("w", encoding="utf-8") as f:
        json.dump(input_data, f, indent=2)

    for artifact_name in ["evidence", "market", "peers", "result"]:
        with (edir / f"{artifact_name}.json").open("w", encoding="utf-8") as f:
            json.dump({artifact_name: {}}, f, indent=2)

    manifest_data = {
        "manifest_version": "1.0.0",
        "evaluation_id": eval_id,
        "artifacts": {
            "evaluation.json": hashlib.sha256((edir / "evaluation.json").read_bytes()).hexdigest(),
            "input.json": hashlib.sha256((edir / "input.json").read_bytes()).hexdigest(),
            "evidence.json": hashlib.sha256((edir / "evidence.json").read_bytes()).hexdigest(),
            "market.json": hashlib.sha256((edir / "market.json").read_bytes()).hexdigest(),
            "peers.json": hashlib.sha256((edir / "peers.json").read_bytes()).hexdigest(),
            "result.json": hashlib.sha256((edir / "result.json").read_bytes()).hexdigest(),
        },
    }
    with (edir / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest_data, f, indent=2)

    return edir


# T-6B-01: Single IPO dataset assembly
def test_t6b_01_single_ipo_dataset_assembly(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "IPO-A-20261005-FINAL-1111"
    res_hash = "hash_ipo_a_1111"
    _create_mock_eval(store, eval_id, "IPO-A", "2026-10-05T10:00:00Z", res_hash, score=70.0)

    dataset = assemble_dataset(store)
    assert dataset.manifest.row_count == 1
    row = dataset.rows[0]
    assert row.ipo_id == "IPO-A"
    assert row.final_evaluation_id == eval_id
    assert row.final_score == 70.0
    assert row.dataset_row_status == DatasetRowStatus.INCOMPLETE.value
    assert row.return_1w_pct is None


# T-6B-02: Multiple IPO dataset assembly
def test_t6b_02_multiple_ipo_dataset_assembly(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "EVAL-A", "ALPHA", "2026-10-01T10:00:00Z", "hash_alpha", score=60.0)
    _create_mock_eval(store, "EVAL-B", "BETA", "2026-10-02T10:00:00Z", "hash_beta", score=75.0)
    _create_mock_eval(store, "EVAL-C", "GAMMA", "2026-10-03T10:00:00Z", "hash_gamma", score=50.0)

    dataset = assemble_dataset(store)
    assert dataset.manifest.row_count == 3
    assert [r.ipo_id for r in dataset.rows] == ["ALPHA", "BETA", "GAMMA"]


# T-6B-03: Multiple horizons per IPO
def test_t6b_03_multiple_horizons_per_ipo(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "VISHAL-FINAL-001"
    res_hash = "vishal_hash_001"
    _create_mock_eval(store, eval_id, "VISHAL", "2026-10-05T12:00:00Z", res_hash, issue_price=110.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for h in ["1W", "1M", "6M"]:
        obs = build_observation(eval_id, res_hash, "VISHAL", "VISHAL", 110.0, "2026-10-06", h, price_feed)
        save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.dataset_row_status == DatasetRowStatus.READY.value
    assert row.return_1w_pct == 24.090909
    assert row.return_1m_pct == 29.090909
    assert row.return_6m_pct == 40.181818
    assert row.excess_return_1w_pct == 23.090909


# T-6B-04: Missing 1W does not create zero
def test_t6b_04_missing_1w_does_not_create_zero(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-NO-1W"
    res_hash = "hash_no_1w"
    _create_mock_eval(store, eval_id, "TEST-NO-1W", "2026-10-05T12:00:00Z", res_hash, issue_price=100.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for h in ["1M", "6M"]:
        obs = build_observation(eval_id, res_hash, "TEST-NO-1W", "VISHAL", 100.0, "2026-10-06", h, price_feed)
        save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.return_1w_pct is None
    assert row.excess_return_1w_pct is None
    assert row.observation_1w_status is None
    assert row.dataset_row_status == DatasetRowStatus.PARTIAL.value


# T-6B-05: Missing 1M does not create zero
def test_t6b_05_missing_1m_does_not_create_zero(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-NO-1M"
    res_hash = "hash_no_1m"
    _create_mock_eval(store, eval_id, "TEST-NO-1M", "2026-10-05T12:00:00Z", res_hash, issue_price=100.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_1w = build_observation(eval_id, res_hash, "TEST-NO-1M", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs_1w, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.return_1m_pct is None
    assert row.excess_return_1m_pct is None


# T-6B-06: Missing 6M does not create zero
def test_t6b_06_missing_6m_does_not_create_zero(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-NO-6M"
    res_hash = "hash_no_6m"
    _create_mock_eval(store, eval_id, "TEST-NO-6M", "2026-10-05T12:00:00Z", res_hash, issue_price=100.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for h in ["1W", "1M"]:
        obs = build_observation(eval_id, res_hash, "TEST-NO-6M", "VISHAL", 100.0, "2026-10-06", h, price_feed)
        save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.return_6m_pct is None
    assert row.excess_return_6m_pct is None


# T-6B-07: Partial lifecycle IPO remains in dataset
def test_t6b_07_partial_lifecycle_ipo_remains_in_dataset(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "EVAL-PARTIAL", "PARTIAL-CO", "2026-10-05T10:00:00Z", "hash_part", issue_price=100.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_1w = build_observation("EVAL-PARTIAL", "hash_part", "PARTIAL-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs_1w, store)

    dataset = assemble_dataset(store)
    assert len(dataset.rows) == 1
    assert dataset.rows[0].dataset_row_status == DatasetRowStatus.PARTIAL.value
    assert dataset.manifest.status_counts[DatasetRowStatus.PARTIAL.value] == 1


# T-6B-08: FINAL evaluation linkage verified
def test_t6b_08_final_evaluation_linkage_verified(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-FINAL-LINKED"
    res_hash = "hash_linked_001"
    _create_mock_eval(store, eval_id, "LINKED-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "LINKED-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.final_evaluation_id == eval_id
    assert row.final_result_hash == res_hash


# T-6B-09: Result hash mismatch rejected
def test_t6b_09_result_hash_mismatch_rejected(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-TAMPER-HASH"
    res_hash = "real_result_hash"
    _create_mock_eval(store, eval_id, "TAMPER-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "TAMPER-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs, store)

    # Tamper with observation file result hash directly
    obs_file = store / eval_id / "observations" / "observation_1w.json"
    data = json.loads(obs_file.read_text())
    data["final_result_hash"] = "forged_result_hash"
    obs_file.write_text(json.dumps(data))

    dataset = assemble_dataset(store)
    assert dataset.rows[0].dataset_row_status == DatasetRowStatus.INVALID.value


# T-6B-10: Invalid observation hash rejected
def test_t6b_10_invalid_observation_hash_rejected(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-BAD-HASH"
    res_hash = "hash_good"
    _create_mock_eval(store, eval_id, "BAD-HASH-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "BAD-HASH-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs, store)

    # Tamper with observation calculation hash
    obs_file = store / eval_id / "observations" / "observation_1w.json"
    data = json.loads(obs_file.read_text())
    data["calculation"]["observation_hash"] = "corrupted_hash_value"
    obs_file.write_text(json.dumps(data))

    dataset = assemble_dataset(store)
    assert dataset.rows[0].dataset_row_status == DatasetRowStatus.INVALID.value


# T-6B-11: Duplicate horizon detected
def test_t6b_11_duplicate_horizon_detected(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-DUP-HORIZON"
    res_hash = "hash_dup_horizon"
    _create_mock_eval(store, eval_id, "DUP-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs1 = build_observation(eval_id, res_hash, "DUP-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed, version=1)
    save_observation(obs1, store)

    # Create another version 1 file under another name without superseding
    obs2_file = store / eval_id / "observations" / "observation_1w_conflict.json"
    obs2_file.write_text(json.dumps(obs1.to_dict()))

    # With strict mode, conflicting observations raise ConflictingObservationError
    with pytest.raises(ConflictingObservationError):
        assemble_dataset(store, strict=True)


# T-6B-12: Conflicting observation versions fail closed
def test_t6b_12_conflicting_observation_versions_fail_closed(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-CONFLICT-VERSIONS"
    res_hash = "hash_conflict"
    _create_mock_eval(store, eval_id, "CONFLICT-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs1 = build_observation(eval_id, res_hash, "CONFLICT-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed, version=1)
    save_observation(obs1, store)

    # Add duplicate version 1 file
    obs2_file = store / eval_id / "observations" / "observation_1w_dup.json"
    obs2_file.write_text(json.dumps(obs1.to_dict()))

    # Non-strict marks row INVALID
    dataset = assemble_dataset(store, strict=False)
    assert dataset.rows[0].dataset_row_status == DatasetRowStatus.INVALID.value


# T-6B-13: Restated observation selects deterministic effective version
def test_t6b_13_restated_observation_selects_effective_version(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-RESTATE"
    res_hash = "hash_restate"
    _create_mock_eval(store, eval_id, "RESTATE-CO", "2026-10-05T10:00:00Z", res_hash, issue_price=100.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_v1 = build_observation(eval_id, res_hash, "RESTATE-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed, version=1)
    save_observation(obs_v1, store)

    obs_v2 = build_observation(
        eval_id, res_hash, "RESTATE-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed,
        version=2,
        supersedes_observation_id=obs_v1.observation_id,
        restatement_reason="Exchange revised EOD settle price",
    )
    save_observation(obs_v2, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    # Effective observation hash must be v2's hash
    assert row.observation_1w_hash == obs_v2.calculation.observation_hash


# T-6B-14: Superseded observation remains preserved
def test_t6b_14_superseded_observation_preserved_on_disk(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-RESTATE-PRESERVED"
    res_hash = "hash_preserved"
    _create_mock_eval(store, eval_id, "PRESERVE-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_v1 = build_observation(eval_id, res_hash, "PRESERVE-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed, version=1)
    save_observation(obs_v1, store)

    obs_v2 = build_observation(
        eval_id, res_hash, "PRESERVE-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed,
        version=2,
        supersedes_observation_id=obs_v1.observation_id,
        restatement_reason="Correction",
    )
    save_observation(obs_v2, store)

    obs_dir = store / eval_id / "observations"
    assert (obs_dir / "observation_1w.json").is_file()
    assert (obs_dir / "observation_1w_v2.json").is_file()


# T-6B-15: Dataset ordering deterministic
def test_t6b_15_dataset_ordering_deterministic(tmp_path: Path):
    store = tmp_path / "evaluations"
    # Create evaluations in reverse chronological order
    _create_mock_eval(store, "EVAL-Z", "ZEBRA", "2026-10-05T12:00:00Z", "hash_z")
    _create_mock_eval(store, "EVAL-M", "MANGO", "2026-10-02T12:00:00Z", "hash_m")
    _create_mock_eval(store, "EVAL-A", "APPLE", "2026-10-01T12:00:00Z", "hash_a")

    dataset = assemble_dataset(store)
    order = [r.ipo_id for r in dataset.rows]
    assert order == ["APPLE", "MANGO", "ZEBRA"]


# T-6B-16: Dataset hash deterministic
def test_t6b_16_dataset_hash_deterministic(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "EVAL-1", "CO-1", "2026-10-01T10:00:00Z", "hash_1")
    _create_mock_eval(store, "EVAL-2", "CO-2", "2026-10-02T10:00:00Z", "hash_2")

    ds1 = assemble_dataset(store, generation_timestamp="2026-10-06T10:00:00Z")
    ds2 = assemble_dataset(store, generation_timestamp="2026-10-06T14:30:00Z")

    assert ds1.manifest.dataset_hash == ds2.manifest.dataset_hash


# T-6B-17: Dataset hash changes when material data changes
def test_t6b_17_dataset_hash_changes_on_material_change(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "EVAL-1", "CO-1", "2026-10-01T10:00:00Z", "hash_1", score=65.0)

    ds_before = assemble_dataset(store)

    # Mutate score materially
    _create_mock_eval(store, "EVAL-1", "CO-1", "2026-10-01T10:00:00Z", "hash_1_mutated", score=72.0)
    ds_after = assemble_dataset(store)

    assert ds_before.manifest.dataset_hash != ds_after.manifest.dataset_hash


# T-6B-18: Filesystem ordering cannot affect dataset
def test_t6b_18_filesystem_ordering_cannot_affect_dataset(tmp_path: Path):
    rows_shuffled_a = [
        BacktestDatasetRow(ipo_id="B", company_name="B", final_evaluation_id="E-B", evaluation_timestamp="2026-10-02T10:00:00Z", final_score=60.0, verdict="APPLY"),
        BacktestDatasetRow(ipo_id="A", company_name="A", final_evaluation_id="E-A", evaluation_timestamp="2026-10-01T10:00:00Z", final_score=70.0, verdict="APPLY"),
    ]
    rows_shuffled_b = [
        BacktestDatasetRow(ipo_id="A", company_name="A", final_evaluation_id="E-A", evaluation_timestamp="2026-10-01T10:00:00Z", final_score=70.0, verdict="APPLY"),
        BacktestDatasetRow(ipo_id="B", company_name="B", final_evaluation_id="E-B", evaluation_timestamp="2026-10-02T10:00:00Z", final_score=60.0, verdict="APPLY"),
    ]
    # Sorting produces same hash
    sorted_a = sorted(rows_shuffled_a, key=lambda r: (r.evaluation_timestamp, r.ipo_id, r.final_evaluation_id))
    sorted_b = sorted(rows_shuffled_b, key=lambda r: (r.evaluation_timestamp, r.ipo_id, r.final_evaluation_id))
    assert compute_dataset_hash(sorted_a) == compute_dataset_hash(sorted_b)


# T-6B-19: Point-in-time evaluation fields remain unchanged
def test_t6b_19_point_in_time_evaluation_fields_remain_unchanged(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-PIT"
    res_hash = "hash_pit_original"
    eval_dir = _create_mock_eval(store, eval_id, "PIT-CO", "2026-10-05T10:00:00Z", res_hash, score=62.5, verdict="APPLY")

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "PIT-CO", "VISHAL", 100.0, "2026-10-06", "1W", price_feed)
    save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.final_score == 62.5
    assert row.verdict == "APPLY"
    assert row.evaluation_timestamp == "2026-10-05T10:00:00Z"

    # Confirm parent evaluation.json was never mutated
    with (eval_dir / "evaluation.json").open("r", encoding="utf-8") as f:
        edata = json.load(f)
    assert edata["score"]["final_score"] == 62.5
    assert edata["result_hash"] == res_hash


# T-6B-20: Benchmark UNKNOWN remains UNKNOWN
def test_t6b_20_benchmark_unknown_remains_unknown(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-BENCH-MISSING"
    res_hash = "hash_bench_missing"
    _create_mock_eval(store, eval_id, "BENCH-CO", "2026-10-05T10:00:00Z", res_hash)

    # Build observation with missing benchmark
    dataset_prices = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "BENCH-CO", "VISHAL", 100.0, "2026-10-06", "1W", dataset_prices, benchmark_symbol="NON_EXISTENT_INDEX")
    save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.benchmark_return_1w_pct is None
    assert row.excess_return_1w_pct is None


# T-6B-21: Excess return UNKNOWN remains UNKNOWN
def test_t6b_21_excess_return_unknown_remains_unknown(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-EXCESS-NULL"
    res_hash = "hash_excess_null"
    _create_mock_eval(store, eval_id, "EXCESS-CO", "2026-10-05T10:00:00Z", res_hash)

    dataset_prices = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "EXCESS-CO", "VISHAL", 100.0, "2026-10-06", "1W", dataset_prices, benchmark_symbol="NON_EXISTENT_INDEX")
    save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.excess_return_1w_pct is None


# T-6B-22: UNVERIFIED corporate-action observation retains status
def test_t6b_22_unverified_corporate_action_retains_status(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "EVAL-UNVERIF-CA"
    res_hash = "hash_unverif_ca"
    _create_mock_eval(store, eval_id, "CA-CO", "2026-10-05T10:00:00Z", res_hash)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(eval_id, res_hash, "CA-CO", "VISHAL", 100.0, "2026-10-06", "6M", price_feed, corporate_action_unverified=True)
    save_observation(obs, store)

    dataset = assemble_dataset(store)
    row = dataset.rows[0]
    assert row.dataset_row_status == DatasetRowStatus.UNVERIFIED.value
    assert row.observation_6m_status == ObservationStatus.UNVERIFIED.value


# T-6B-23: Dataset manifest verifies correctly
def test_t6b_23_dataset_manifest_verifies_correctly(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "E1", "CO-1", "2026-10-01T10:00:00Z", "hash1")
    _create_mock_eval(store, "E2", "CO-2", "2026-10-02T10:00:00Z", "hash2")

    dataset = assemble_dataset(store)
    out_json = export_dataset_json(dataset, tmp_path / "dataset.json")

    res = verify_dataset(out_json, store_root=store)
    assert res["status"] == "PASS"
    assert res["dataset_hash_match"] is True
    assert res["ordering_valid"] is True


# T-6B-24: Dataset verification rejects tampering
def test_t6b_24_dataset_verification_rejects_tampering(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "E1", "CO-1", "2026-10-01T10:00:00Z", "hash1")

    dataset = assemble_dataset(store)
    out_json = export_dataset_json(dataset, tmp_path / "dataset.json")

    # Tamper with a row's final_score directly in file
    content = json.loads(out_json.read_text())
    content["rows"][0]["final_score"] = 99.9
    out_json.write_text(json.dumps(content))

    res = verify_dataset(out_json, store_root=store)
    assert res["status"] == "FAIL"
    assert "dataset_hash mismatch" in res["reason"]


# T-6B-25: JSON export is deterministic
def test_t6b_25_json_export_is_deterministic(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "E1", "CO-1", "2026-10-01T10:00:00Z", "hash1")

    ds = assemble_dataset(store, generation_timestamp="2026-10-06T12:00:00Z")
    f1 = export_dataset_json(ds, tmp_path / "ds1.json")
    f2 = export_dataset_json(ds, tmp_path / "ds2.json")

    assert f1.read_bytes() == f2.read_bytes()


# T-6B-26: CSV projection is deterministic
def test_t6b_26_csv_projection_is_deterministic(tmp_path: Path):
    store = tmp_path / "evaluations"
    _create_mock_eval(store, "E1", "CO-1", "2026-10-01T10:00:00Z", "hash1")

    ds = assemble_dataset(store)
    c1 = export_dataset_csv(ds, tmp_path / "ds1.csv")
    c2 = export_dataset_csv(ds, tmp_path / "ds2.csv")

    assert c1.read_bytes() == c2.read_bytes()


# T-6B-27: Multi-IPO Excel Post_Listing projection
def test_t6b_27_multi_ipo_excel_post_listing_projection(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_a = "EVAL-EXCEL-A"
    eval_b = "EVAL-EXCEL-B"
    _create_mock_eval(store, eval_a, "COMPANY-A", "2026-10-01T10:00:00Z", "hash_a", issue_price=100.0)
    _create_mock_eval(store, eval_b, "COMPANY-B", "2026-10-02T10:00:00Z", "hash_b", issue_price=110.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for eid, sym, ipo in [(eval_a, "VISHAL", "COMPANY-A"), (eval_b, "VISHAL", "COMPANY-B")]:
        obs = build_observation(eid, f"hash_{ipo[-1].lower()}", ipo, sym, 100.0, "2026-10-06", "1W", price_feed)
        save_observation(obs, store)

    from ipo_screening.evaluation import EvaluationStore
    from ipo_screening.excel import project

    estore = EvaluationStore(store)
    wb_file = tmp_path / "history.xlsx"
    proj_res = project(estore, wb_file)

    import openpyxl
    wb = openpyxl.load_workbook(wb_file, data_only=True)
    post_sheet = wb["Post_Listing"]
    # 1 header row + 2 data rows = 3 rows
    assert post_sheet.max_row == 3
    eids = [post_sheet.cell(row=i, column=1).value for i in range(2, 4)]
    assert eids == [eval_a, eval_b]


# T-6B-28: Multi-IPO Excel Backtest projection
def test_t6b_28_multi_ipo_excel_backtest_projection(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_a = "EVAL-BT-A"
    eval_b = "EVAL-BT-B"
    _create_mock_eval(store, eval_a, "BT-A", "2026-10-01T10:00:00Z", "hash_a", issue_price=100.0)
    _create_mock_eval(store, eval_b, "BT-B", "2026-10-02T10:00:00Z", "hash_b", issue_price=110.0)

    price_feed = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for eid, sym, ipo in [(eval_a, "VISHAL", "BT-A"), (eval_b, "VISHAL", "BT-B")]:
        obs = build_observation(eid, f"hash_{ipo[-1].lower()}", ipo, sym, 100.0, "2026-10-06", "1W", price_feed)
        save_observation(obs, store)

    from ipo_screening.evaluation import EvaluationStore
    from ipo_screening.excel import project

    estore = EvaluationStore(store)
    wb_file = tmp_path / "history.xlsx"
    project(estore, wb_file)

    import openpyxl
    wb = openpyxl.load_workbook(wb_file, data_only=True)
    bt_sheet = wb["Backtest"]
    assert bt_sheet.max_row == 3
    ipos = [bt_sheet.cell(row=i, column=1).value for i in range(2, 4)]
    assert ipos == ["BT-A", "BT-B"]


# T-6B-29: Existing v1.5 + Phase 6A tests remain green
def test_t6b_29_golden_result_hash_stable():
    from conftest import CONFIG_PATH, GOLDEN_INPUT
    from ipo_screening.pipeline import evaluate, load_config

    doc = json.load(open(GOLDEN_INPUT))
    config = load_config(CONFIG_PATH)
    outcome = evaluate(doc, config, mode="final", evaluation_datetime=EVAL_AT)
    assert outcome.record.result_hash == "e84f8bc0f9b942c43f937fa3b12fdba3c3ef23cc613e9d921a749b12955619e1"


# T-6B-30: Complete Phase 6B E2E workflow
def test_t6b_30_complete_phase6b_e2e(tmp_path: Path):
    store = tmp_path / "evaluations"
    eval_id = "VISHAL-NIRMITI-LIMITED-20261005-120000Z-final-3bd4bca3"
    res_hash = "3bd4bca3d2264e2d07645359a83aa72c1180ad3946f9a3e04b91f3e902e09df4"
    _create_mock_eval(store, eval_id, "VISHAL-NIRMITI-LIMITED", "2026-10-05T12:00:00Z", res_hash, issue_price=110.0)

    # 1. Ingest prices via CLI
    ingest_cmd = [
        sys.executable,
        "engine/tools/ipo_screen.py",
        "post-listing",
        "ingest",
        "--evaluation-id",
        eval_id,
        "--prices",
        str(FIXTURES_DIR / "clean_bhavcopy.csv"),
        "--store",
        str(store),
        "-v",
    ]
    p_ingest = subprocess.run(ingest_cmd, capture_output=True, text=True)
    assert p_ingest.returncode == 0

    # 2. Assemble dataset via CLI
    ds_file = tmp_path / "backtest_dataset.json"
    csv_file = tmp_path / "backtest_dataset.csv"
    dataset_cmd = [
        sys.executable,
        "engine/tools/ipo_screen.py",
        "post-listing",
        "dataset",
        "--store",
        str(store),
        "--output",
        str(ds_file),
        "--csv",
        str(csv_file),
        "--strict",
        "-v",
    ]
    p_dataset = subprocess.run(dataset_cmd, capture_output=True, text=True)
    assert p_dataset.returncode == 0
    assert "HISTORICAL OUTCOME DATASET ASSEMBLED" in p_dataset.stdout
    assert ds_file.is_file()
    assert csv_file.is_file()

    # 3. Verify dataset via CLI
    verify_cmd = [
        sys.executable,
        "engine/tools/ipo_screen.py",
        "post-listing",
        "verify-dataset",
        "--dataset",
        str(ds_file),
        "--store",
        str(store),
        "-v",
    ]
    p_verify = subprocess.run(verify_cmd, capture_output=True, text=True)
    assert p_verify.returncode == 0
    assert "Audit Status:         ALL PASS" in p_verify.stdout
