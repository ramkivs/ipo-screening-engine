"""Phase 6A Acceptance Test Suite: Post-Listing Observation Model & Deterministic Return Engine.

Covers tests T-6A-01 through T-6A-24:
- Immutable observation data model and schema validation
- Linked storage under <store>/<final-evaluation-id>/observations/
- Calendar horizon expansion and trading day clamping
- Deterministic listing gain, absolute return, benchmark return, and excess return
- Corporate actions handling and verification flags
- Strict fail-closed behavior for negative, malformed, or missing inputs
- Deterministic observation hashing and audit verification
- CLI subcommand `post-listing ingest`
- Backward-compatible Excel projection
"""

from __future__ import annotations

import dataclasses
import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import Any, Dict

import pytest

try:
    from ipo_screening.post_listing import (
        CALCULATION_VERSION,
        BenchmarkObservation,
        CalculationMetadata,
        DuplicatePriceConflictError,
        EvaluationNotFoundError,
        Horizon,
        MalformedPriceDataError,
        NegativePriceError,
        NotFinalEvaluationError,
        ObservationStatus,
        PostListingObservation,
        PriceObservation,
        ProvenanceRecord,
        ResultHashMismatchError,
        ReturnSet,
        VerificationStatus,
        build_observation,
        calculate_absolute_return,
        calculate_benchmark_return,
        calculate_excess_return,
        calculate_listing_gain,
        calculate_secondary_return,
        compute_observation_hashes,
        compute_target_date,
        format_iso_date,
        get_observations_dir,
        is_trading_day,
        list_observations,
        load_price_file,
        parse_iso_date,
        previous_trading_day,
        read_observation,
        resolve_observation_date,
        save_observation,
        sha256_canonical_dict,
        validate_parent_evaluation,
        verify_observation_hashes,
    )
except ImportError:
    from engine.ipo_screening.post_listing import (
        CALCULATION_VERSION,
        BenchmarkObservation,
        CalculationMetadata,
        DuplicatePriceConflictError,
        EvaluationNotFoundError,
        Horizon,
        MalformedPriceDataError,
        NegativePriceError,
        NotFinalEvaluationError,
        ObservationStatus,
        PostListingObservation,
        PriceObservation,
        ProvenanceRecord,
        ResultHashMismatchError,
        ReturnSet,
        VerificationStatus,
        build_observation,
        calculate_absolute_return,
        calculate_benchmark_return,
        calculate_excess_return,
        calculate_listing_gain,
        calculate_secondary_return,
        compute_observation_hashes,
        compute_target_date,
        format_iso_date,
        get_observations_dir,
        is_trading_day,
        list_observations,
        load_price_file,
        parse_iso_date,
        previous_trading_day,
        read_observation,
        resolve_observation_date,
        save_observation,
        sha256_canonical_dict,
        validate_parent_evaluation,
        verify_observation_hashes,
    )

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "post_listing"


def _create_mock_final_evaluation(store_dir: Path, eval_id: str, result_hash: str, ipo_id: str = "VISHAL", issue_price: float = 110.0, listing_date: str = "2026-10-06") -> Path:
    eval_dir = store_dir / eval_id
    eval_dir.mkdir(parents=True, exist_ok=True)
    eval_payload = {
        "evaluation_id": eval_id,
        "ipo_id": ipo_id,
        "evaluation_mode": "FINAL",
        "evaluation_timestamp": "2026-10-06T10:00:00Z",
        "result_hash": result_hash,
        "score": {"final_score": 68.5},
        "verdict": {"verdict": "APPLY"},
        "knockouts": {"status": "PASS"},
    }
    with (eval_dir / "evaluation.json").open("w", encoding="utf-8") as f:
        json.dump(eval_payload, f, indent=2)

    input_payload = {
        "snapshot": {
            "input": {
                "ipo_id": ipo_id,
                "company": {"name": "Vishal Nirmiti Ltd", "symbol": "VISHAL"},
                "issue": {"issue_price": issue_price, "listing_date": listing_date},
            }
        }
    }
    with (eval_dir / "input.json").open("w", encoding="utf-8") as f:
        json.dump(input_payload, f, indent=2)

    for artifact_name in ["evidence", "market", "peers", "result"]:
        with (eval_dir / f"{artifact_name}.json").open("w", encoding="utf-8") as f:
            json.dump({artifact_name: {}}, f, indent=2)

    manifest_payload = {
        "manifest_version": "1.0.0",
        "evaluation_id": eval_id,
        "artifacts": {
            "evaluation.json": hashlib.sha256((eval_dir / "evaluation.json").read_bytes()).hexdigest(),
            "input.json": hashlib.sha256((eval_dir / "input.json").read_bytes()).hexdigest(),
            "evidence.json": hashlib.sha256((eval_dir / "evidence.json").read_bytes()).hexdigest(),
            "market.json": hashlib.sha256((eval_dir / "market.json").read_bytes()).hexdigest(),
            "peers.json": hashlib.sha256((eval_dir / "peers.json").read_bytes()).hexdigest(),
            "result.json": hashlib.sha256((eval_dir / "result.json").read_bytes()).hexdigest(),
        }
    }
    with (eval_dir / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest_payload, f, indent=2)

    return eval_dir


# T-6A-01: Canonical observation structure conforms to specification
def test_t6a_01_canonical_observation_structure():
    prices = PriceObservation(issue_price=110.0, listing_open=125.0, raw_observed_close=136.5, corporate_action_factor=1.0, adjusted_observed_close=136.5)
    bench = BenchmarkObservation(symbol="NIFTY_50_TRI", raw_listing_value=25050.0, raw_observed_value=25300.5, return_pct=1.0)
    rets = ReturnSet(listing_gain_pct=13.636364, absolute_return_pct=24.090909, secondary_return_pct=9.2, benchmark_return_pct=1.0, excess_return_pct=23.090909)
    prov = ProvenanceRecord(source_id="BHAV-001", source_type="CSV_BHAVCOPY", source_uri_or_file="bhav.csv", retrieval_timestamp="2026-10-06T12:00:00Z", content_hash="abc123")
    calc = CalculationMetadata(calculation_version="1.0.0", calculation_inputs_hash="hash1", observation_hash="hash2")

    obs = PostListingObservation(
        observation_id="OBS-EVAL1-1W",
        final_evaluation_id="EVAL1",
        final_result_hash="resulthash123",
        ipo_id="VISHAL",
        horizon="1W",
        listing_date="2026-10-06",
        target_observation_date="2026-10-13",
        actual_observation_date="2026-10-13",
        prices=prices,
        benchmark=bench,
        returns=rets,
        provenance=prov,
        calculation=calc,
    )
    d = obs.to_dict()
    assert d["observation_id"] == "OBS-EVAL1-1W"
    assert d["final_evaluation_id"] == "EVAL1"
    assert d["horizon"] == "1W"
    assert d["prices"]["issue_price"] == 110.0
    assert d["returns"]["excess_return_pct"] == 23.090909
    assert d["observation_status"] == ObservationStatus.VERIFIED.value


# T-6A-02: Observation schema validation (all required fields, typing, enums)
def test_t6a_02_observation_roundtrip_dict():
    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id="EVAL-001",
        final_result_hash="abc111",
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
    )
    d = obs.to_dict()
    restored = PostListingObservation.from_dict(d)
    assert restored.observation_id == obs.observation_id
    assert restored.prices.issue_price == obs.prices.issue_price
    assert restored.calculation.observation_hash == obs.calculation.observation_hash
    assert restored.returns.absolute_return_pct == obs.returns.absolute_return_pct


# T-6A-03: Observation immutability (frozen dataclass mutation raises error)
def test_t6a_03_observation_immutability():
    prices = PriceObservation(issue_price=100.0)
    with pytest.raises(dataclasses.FrozenInstanceError):
        prices.issue_price = 105.0  # type: ignore

    obs = PostListingObservation(
        observation_id="OBS-1",
        final_evaluation_id="E1",
        final_result_hash="R1",
        ipo_id="VISHAL",
        horizon="1W",
        listing_date="2026-10-06",
        target_observation_date="2026-10-13",
        actual_observation_date="2026-10-13",
        prices=prices,
        benchmark=BenchmarkObservation(symbol="NIFTY_50_TRI"),
        returns=ReturnSet(),
        provenance=ProvenanceRecord("s1", "csv", "f.csv", "t", "h"),
        calculation=CalculationMetadata(),
    )
    with pytest.raises(dataclasses.FrozenInstanceError):
        obs.horizon = "1M"  # type: ignore


# T-6A-04: Observation storage location follows <store>/<final-evaluation-id>/observations/
def test_t6a_04_observation_storage_location(tmp_path: Path):
    eval_id = "VISHAL-FINAL-001"
    res_hash = "9999aaaa8888bbbb"
    _create_mock_final_evaluation(tmp_path, eval_id, res_hash)

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash=res_hash,
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
    )
    saved_file = save_observation(obs, tmp_path)
    expected_path = tmp_path / eval_id / "observations" / "observation_1w.json"
    assert saved_file == expected_path
    assert saved_file.is_file()


# T-6A-05: Storage refuses attachment to preliminary evaluation
def test_t6a_05_storage_refuses_preliminary_evaluation(tmp_path: Path):
    eval_id = "VISHAL-PRELIM-001"
    eval_dir = tmp_path / eval_id
    eval_dir.mkdir(parents=True, exist_ok=True)
    with (eval_dir / "evaluation.json").open("w") as f:
        json.dump({"evaluation_id": eval_id, "evaluation_mode": "PRELIMINARY", "result_hash": "prelimhash"}, f)

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash="prelimhash",
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
    )
    with pytest.raises(NotFinalEvaluationError, match="must be FINAL"):
        save_observation(obs, tmp_path)


# T-6A-06: Storage refuses attachment when evaluation result_hash mismatches
def test_t6a_06_storage_refuses_result_hash_mismatch(tmp_path: Path):
    eval_id = "VISHAL-FINAL-002"
    _create_mock_final_evaluation(tmp_path, eval_id, "real_hash_123")

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash="forged_hash_999",
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
    )
    with pytest.raises(ResultHashMismatchError):
        save_observation(obs, tmp_path)


# T-6A-07: Storage never overwrites or mutates parent evaluation artifacts
def test_t6a_07_parent_evaluation_artifacts_untouched(tmp_path: Path):
    eval_id = "VISHAL-FINAL-003"
    res_hash = "hash_frozen_003"
    eval_dir = _create_mock_final_evaluation(tmp_path, eval_id, res_hash)

    eval_json_before = (eval_dir / "evaluation.json").read_bytes()
    input_json_before = (eval_dir / "input.json").read_bytes()
    manifest_before = (eval_dir / "manifest.json").read_bytes()

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    for h in ["1W", "1M", "6M"]:
        obs = build_observation(
            final_evaluation_id=eval_id,
            final_result_hash=res_hash,
            ipo_id="VISHAL",
            symbol="VISHAL",
            issue_price=110.0,
            listing_date="2026-10-06",
            horizon=h,
            price_dataset=dataset,
        )
        save_observation(obs, tmp_path)

    assert (eval_dir / "evaluation.json").read_bytes() == eval_json_before
    assert (eval_dir / "input.json").read_bytes() == input_json_before
    assert (eval_dir / "manifest.json").read_bytes() == manifest_before


# T-6A-08: Observation manifest records correct SHA-256 for each observation file
def test_t6a_08_observation_manifest_hashes(tmp_path: Path):
    eval_id = "VISHAL-FINAL-004"
    res_hash = "hash_frozen_004"
    eval_dir = _create_mock_final_evaluation(tmp_path, eval_id, res_hash)

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash=res_hash,
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
    )
    save_observation(obs, tmp_path)

    obs_manifest_file = eval_dir / "observations" / "manifest.json"
    assert obs_manifest_file.is_file()
    with obs_manifest_file.open("r", encoding="utf-8") as f:
        mdata = json.load(f)

    obs_entry = mdata["observations"]["observation_1w.json"]
    actual_file_sha = hashlib.sha256((eval_dir / "observations" / "observation_1w.json").read_bytes()).hexdigest()
    assert obs_entry["file_sha256"] == actual_file_sha

    audit = verify_observation_hashes(tmp_path, eval_id)
    assert audit["status"] == "OK"
    assert len(audit["items"]) == 1
    assert audit["items"][0]["status"] == "OK"


# T-6A-09: Restatement creates new observation version preserving original
def test_t6a_09_restatement_preserves_original(tmp_path: Path):
    eval_id = "VISHAL-FINAL-005"
    res_hash = "hash_frozen_005"
    eval_dir = _create_mock_final_evaluation(tmp_path, eval_id, res_hash)

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_v1 = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash=res_hash,
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
        version=1,
    )
    save_observation(obs_v1, tmp_path)
    v1_bytes = (eval_dir / "observations" / "observation_1w.json").read_bytes()

    obs_v2 = build_observation(
        final_evaluation_id=eval_id,
        final_result_hash=res_hash,
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="1W",
        price_dataset=dataset,
        version=2,
        supersedes_observation_id=obs_v1.observation_id,
        restatement_reason="Exchange revised EOD close for block trade adjustment",
    )
    save_observation(obs_v2, tmp_path)

    assert (eval_dir / "observations" / "observation_1w.json").read_bytes() == v1_bytes
    assert (eval_dir / "observations" / "observation_1w_v2.json").is_file()


# T-6A-10: Restatement links supersedes_observation_id and records restatement_reason
def test_t6a_10_restatement_linkage_and_audit(tmp_path: Path):
    eval_id = "VISHAL-FINAL-006"
    res_hash = "hash_frozen_006"
    _create_mock_final_evaluation(tmp_path, eval_id, res_hash)

    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs_v1 = build_observation(eval_id, res_hash, "VISHAL", "VISHAL", 110.0, "2026-10-06", "6M", dataset, version=1)
    save_observation(obs_v1, tmp_path)

    obs_v2 = build_observation(
        eval_id, res_hash, "VISHAL", "VISHAL", 110.0, "2026-10-06", "6M", dataset,
        version=2,
        supersedes_observation_id=obs_v1.observation_id,
        restatement_reason="Bonus issue adjustment factor applied",
    )
    save_observation(obs_v2, tmp_path)

    read_v2 = read_observation(tmp_path, eval_id, "6M", version=2)
    assert read_v2 is not None
    assert read_v2.supersedes_observation_id == obs_v1.observation_id
    assert read_v2.restatement_reason == "Bonus issue adjustment factor applied"

    all_obs = list_observations(tmp_path, eval_id)
    assert len(all_obs) == 2


# T-6A-11: Calendar resolution 1W = 7 calendar days clamped to trading day
def test_t6a_11_calendar_resolution_1w():
    # 2026-10-06 (Tue) + 7 days = 2026-10-13 (Tue)
    target = compute_target_date("2026-10-06", "1W")
    assert target == "2026-10-13"

    trading_dates = {"2026-10-06", "2026-10-13"}
    t_target, t_actual = resolve_observation_date("2026-10-06", "1W", trading_dates)
    assert t_target == "2026-10-13"
    assert t_actual == "2026-10-13"


# T-6A-12: Calendar resolution 1M = 30 calendar days clamped to trading day
def test_t6a_12_calendar_resolution_1m():
    # 2026-10-06 (Tue) + 30 days = 2026-11-05 (Thu)
    target = compute_target_date("2026-10-06", "1M")
    assert target == "2026-11-05"

    trading_dates = {"2026-10-06", "2026-11-05"}
    t_target, t_actual = resolve_observation_date("2026-10-06", "1M", trading_dates)
    assert t_target == "2026-11-05"
    assert t_actual == "2026-11-05"


# T-6A-13: Calendar resolution 6M = 180 calendar days clamped to trading day
def test_t6a_13_calendar_resolution_6m():
    # 2026-10-06 (Tue) + 180 days = 2027-04-04 (Sun)
    target = compute_target_date("2026-10-06", "6M")
    assert target == "2027-04-04"

    # Sunday 2027-04-04 clamped to preceding Friday 2027-04-02
    trading_dates = {"2026-10-06", "2027-04-02"}
    t_target, t_actual = resolve_observation_date("2026-10-06", "6M", trading_dates)
    assert t_target == "2027-04-04"
    assert t_actual == "2027-04-02"


# T-6A-14: Target date landing on weekend clamps to latest preceding Friday
def test_t6a_14_weekend_clamping():
    # Target date Saturday 2026-10-10 or Sunday 2026-10-11
    # Trading dates include Friday 2026-10-09
    trading_dates = {"2026-10-08", "2026-10-09", "2026-10-12"}
    assert previous_trading_day("2026-10-10", trading_dates) == "2026-10-09"
    assert previous_trading_day("2026-10-11", trading_dates) == "2026-10-09"


# T-6A-15: Target date landing on holiday clamps to latest preceding trading day
def test_t6a_15_holiday_clamping():
    # Sunday 2027-04-04 target date, but Friday 2027-04-02 is a holiday (not in trading dates)
    # Available trading date is Thursday 2027-04-01
    dataset = load_price_file(FIXTURES_DIR / "weekend_holiday_bhavcopy.csv")
    t_target, t_actual = resolve_observation_date("2026-10-06", "6M", dataset.trading_dates)
    assert t_target == "2027-04-04"
    assert t_actual == "2027-04-01"


# T-6A-16: Listing gain calculation matches (listing_open - issue_price) / issue_price * 100
def test_t6a_16_listing_gain_calculation():
    # issue_price = 110.0, listing_open = 125.0
    # (125 - 110) / 110 * 100 = 15 / 110 * 100 = 13.636364%
    gain = calculate_listing_gain(110.0, 125.0)
    assert gain == 13.636364

    # Missing open
    assert calculate_listing_gain(110.0, None) is None

    # Invalid issue price
    with pytest.raises(ValueError, match="positive"):
        calculate_listing_gain(0.0, 125.0)
    with pytest.raises(ValueError, match="positive"):
        calculate_listing_gain(-10.0, 125.0)


# T-6A-17: 1W, 1M, 6M absolute return calculation matches (adj_close - issue_price) / issue_price * 100
def test_t6a_17_absolute_return_calculation():
    # 1W: issue = 110.0, observed_close = 136.5, factor = 1.0 -> (136.5 - 110) / 110 * 100 = 24.090909%
    adj_close, ret = calculate_absolute_return(110.0, 136.5, 1.0)
    assert adj_close == 136.5
    assert ret == 24.090909

    # 1M: issue = 110.0, observed_close = 142.0, factor = 1.0 -> (142 - 110) / 110 * 100 = 29.090909%
    adj_close, ret = calculate_absolute_return(110.0, 142.0, 1.0)
    assert adj_close == 142.0
    assert ret == 29.090909

    # Missing observed close -> None
    adj_close, ret = calculate_absolute_return(110.0, None, 1.0)
    assert adj_close is None
    assert ret is None


# T-6A-18: Benchmark return calculation matches formula
def test_t6a_18_benchmark_return_calculation():
    # listing_bench = 25050.0, observed_bench = 25300.5
    # (25300.5 - 25050.0) / 25050.0 * 100 = 250.5 / 25050.0 * 100 = 1.0%
    b_ret = calculate_benchmark_return(25050.0, 25300.5)
    assert b_ret == 1.0

    # Missing benchmark -> None
    assert calculate_benchmark_return(25050.0, None) is None
    assert calculate_benchmark_return(None, 25300.5) is None


# T-6A-19: Excess return calculation matches absolute_return - benchmark_return
def test_t6a_19_excess_return_calculation():
    # absolute = 24.090909%, benchmark = 1.0% -> excess = 23.090909%
    excess = calculate_excess_return(24.090909, 1.0)
    assert excess == 23.090909


# T-6A-20: Missing benchmark returns fail-closed None / UNKNOWN for excess return, never defaults to 0.0
def test_t6a_20_missing_benchmark_fail_closed_unknown():
    assert calculate_excess_return(24.090909, None) is None
    assert calculate_excess_return(None, 1.0) is None
    assert calculate_excess_return(None, None) is None


# T-6A-21: Corporate action factor adjustment adjusted = raw * factor verified
def test_t6a_21_corporate_action_factor():
    dataset = load_price_file(FIXTURES_DIR / "corporate_action_bhavcopy.csv")
    # For 6M: raw_close = 80.0, corporate_action_factor = 2.0
    # adjusted_close = 160.0. With issue_price = 110.0:
    # return = (160 - 110) / 110 * 100 = 45.454545%
    obs = build_observation(
        final_evaluation_id="EVAL-CA",
        final_result_hash="hash_ca",
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="6M",
        price_dataset=dataset,
    )
    assert obs.prices.raw_observed_close == 80.0
    assert obs.prices.corporate_action_factor == 2.0
    assert obs.prices.adjusted_observed_close == 160.0
    assert obs.returns.absolute_return_pct == 45.454545
    assert obs.verification_status == VerificationStatus.VERIFIED.value


# T-6A-22: Unverified corporate action marks observation status UNVERIFIED
def test_t6a_22_unverified_corporate_action():
    dataset = load_price_file(FIXTURES_DIR / "clean_bhavcopy.csv")
    obs = build_observation(
        final_evaluation_id="EVAL-UNVERIFIED",
        final_result_hash="hash_unverified",
        ipo_id="VISHAL",
        symbol="VISHAL",
        issue_price=110.0,
        listing_date="2026-10-06",
        horizon="6M",
        price_dataset=dataset,
        corporate_action_unverified=True,
    )
    assert obs.verification_status == VerificationStatus.UNVERIFIED.value
    assert obs.observation_status == ObservationStatus.UNVERIFIED.value


# T-6A-23: Malformed, negative, or conflicting prices in price file fail-closed with error
def test_t6a_23_malformed_negative_conflicting_prices():
    # Negative price fails
    with pytest.raises(NegativePriceError):
        load_price_file(FIXTURES_DIR / "corrupt_negative_bhavcopy.csv")

    # Conflicting duplicate prices fail
    with pytest.raises(DuplicatePriceConflictError):
        load_price_file(FIXTURES_DIR / "conflicting_bhavcopy.csv")

    # Missing file fails
    with pytest.raises(FileNotFoundError):
        load_price_file("non_existent_file.csv")


# T-6A-24: Complete end-to-end integration: ipo_screen.py post-listing ingest on stored FINAL evaluation
def test_t6a_24_e2e_cli_post_listing_ingest(tmp_path: Path):
    eval_id = "VISHAL-NIRMITI-LIMITED-20261005-120000Z-final-3bd4bca3"
    res_hash = "3bd4bca3d2264e2d07645359a83aa72c1180ad3946f9a3e04b91f3e902e09df4"
    _create_mock_final_evaluation(tmp_path, eval_id, res_hash, ipo_id="VISHAL-NIRMITI-LIMITED", issue_price=110.0, listing_date="2026-10-06")

    # Run CLI post-listing ingest
    cmd = [
        sys.executable,
        "engine/tools/ipo_screen.py",
        "post-listing",
        "ingest",
        "--evaluation-id",
        eval_id,
        "--prices",
        str(FIXTURES_DIR / "clean_bhavcopy.csv"),
        "--store",
        str(tmp_path),
        "-v",
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    assert proc.returncode == 0, f"CLI ingest failed: {proc.stderr}"
    assert "POST-LISTING INGESTION REPORT" in proc.stdout
    assert "Verification Status:  PASS" in proc.stdout

    # Verify observations created
    obs_dir = tmp_path / eval_id / "observations"
    assert (obs_dir / "observation_1w.json").is_file()
    assert (obs_dir / "observation_1m.json").is_file()
    assert (obs_dir / "observation_6m.json").is_file()
    assert (obs_dir / "manifest.json").is_file()

    # Re-verify audit hashes
    audit = verify_observation_hashes(tmp_path, eval_id)
    assert audit["status"] == "OK"
    assert len(audit["items"]) == 3
    for it in audit["items"]:
        assert it["status"] == "OK"
        assert it["result_hash_match"] is True
        assert it["observation_hash_match"] is True

    # Test projection into Excel workbook
    workbook_path = tmp_path / "IPO_Screening_History.xlsx"
    proj_cmd = [
        sys.executable,
        "engine/tools/ipo_screen.py",
        "project",
        "--store",
        str(tmp_path),
        "--workbook",
        str(workbook_path),
    ]
    p_proc = subprocess.run(proj_cmd, capture_output=True, text=True)
    assert p_proc.returncode == 0, f"CLI project failed: {p_proc.stderr}"
    assert workbook_path.is_file()

    # Read back workbook using openpyxl to confirm Post_Listing and Backtest sheets have rows
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    post_sheet = wb["Post_Listing"]
    backtest_sheet = wb["Backtest"]

    # Header is row 1, data is row 2
    assert post_sheet.max_row >= 2
    assert backtest_sheet.max_row >= 2
    row2_post = [cell.value for cell in post_sheet[2]]
    # Column 0: evaluation_id, Column 4: listing_gain_pct, Column 5: return_1w_pct
    assert row2_post[0] == eval_id
    assert row2_post[4] == 13.636364  # listing gain
    assert row2_post[5] == 24.090909  # 1W return
